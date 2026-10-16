#!/usr/bin/env python3
"""Planilha para a previsão orçamentária dos contratos OGU — grava conferencia/previsao_orcamentaria_OGU.xlsx (fora do git).

Três grupos, iguais aos slides (padrão de cada um):
  - sem edital publicado (3.2, publicação do edital) e edital publicado sem homologação (3.2, fim da licitação):
    OGU em "Contratado - Normal" no painelpac, PC 72 incluída em simulação;
  - obras iniciadas OGU (5.2 com fonte OGU, migradas + novas, situação SACI + balanço), + itens do balanço sem TCI.
Usa a última extração (dados.load()) e, só para leitura, o banco (proposta, operação, tomador, valores, desembolso por ano).

Uso (na raiz do projeto):
  uv run --with pandas --with openpyxl --with python-dotenv --with msoffcrypto-tool --with sqlalchemy --with psycopg2-binary \
         python python/previsao_orcamentaria.py
"""
import os
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import text

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)
sys.path.insert(0, os.path.join(RAIZ, '..', 'painelgm', 'python', 'saci'))
import dados  # noqa: E402
from conexao import criar_engine  # noqa: E402

SAIDA = os.path.join(RAIZ, 'conferencia', 'previsao_orcamentaria_OGU.xlsx')
D = dados.load(); s = D['saci'].copy(); o = D['ogu'].copy(); bx = D['bx']; ext = D['extracao']
REF = pd.Timestamp(ext['extraido_em'][:10]); LIM = pd.Timestamp('2026-10-30'); CORTE = pd.Timestamp('2025-10-21')
dt = lambda x: pd.to_datetime(x, errors='coerce')

# ---------- 155 / 257: régua do 3.2
r = o[(o.dsc_situacao_contrato_mcid_saci == 'Contratado - Normal') & ~o.excl & ~o.semmod].copy()
r['pc72'] = dt(r.dte_assinatura_contrato_saci) < CORTE
lae, pub = dt(r.dte_primeira_data_lae_tdb), dt(r.dte_publicacao_licitacao_tgov)
semed = pub.isna(); semhom = pub.notna() & dt(r.dte_homologacao_licitacao_tgov).isna()
legal_pub = (lae + pd.Timedelta(days=60)).where(r.pc72, dt(r.prazo_pub_licitacao_36a_calc))
legal_hom = (pub + pd.Timedelta(days=120)).where(r.pc72, dt(r.prazo_homolog_licitacao_120d))


def efet(legal, calc, pc72):
    base = legal.where(legal > LIM, LIM).where(legal.notna())
    return base.where(pc72, dt(calc).fillna(base))


r['legal'] = legal_pub.where(semed, legal_hom)
r['efetivo'] = efet(legal_pub, r.prazo_pub_licitacao_calc, r.pc72).where(semed, efet(legal_hom, r.prazo_homolog_licitacao_calc, r.pc72))
g155 = r[semed & legal_pub.notna()].copy(); g257 = r[semhom & legal_hom.notna()].copy()

# ---------- 734: obras iniciadas OGU (5.2)
SITO = {'Não iniciada': 0, 'Não cadastrada': 0, 'Cadastrada': 0, 'Sem informação': 0, 'Em execução': 1, 'Paralisada': 2,
        'Concluída': 3, 'Inacabada': 4, 'Cancelada': 5, 'Inativada': 5}
a = s[~s.excl & ~s.semmod & (s.txt_fonte == 'OGU') & ~s.dsc_situacao_contrato_mcid.isin(['Em Contratação', 'Cancelado ou Distratado'])].copy()
so = a.dsc_situacao_objeto_mcid.map(SITO).fillna(0).astype(int)


def eff(sv, g):  # mesma regra do slide: SACI + balanço
    if g == 0: return sv
    if g == 3: return 3
    if 1 <= sv <= 4: return sv
    return 1 if g == 1 else 2


a['sit_eff'] = [eff(x, y) for x, y in zip(so, a.bal_g)]
g734 = a[a.sit_eff.between(1, 4)].copy()
bxo = bx[(bx.fonte == 'OGU') & (bx['mod'] != dados.OUTRAS)].copy()  # sem modalidade fica fora (exceção provisória)

# ---------- banco (só leitura)
tcis = sorted(set(g155.cod_tci_saci) | set(g257.cod_tci_saci) | set(g734.cod_tci))
eng = criar_engine(connect_args={'connect_timeout': 15, 'options': '-c default_transaction_read_only=on -c statement_timeout=120000'})
with eng.connect() as c:
    db = pd.read_sql(text("""SELECT cod_tci, num_proposta, cod_operacao, txt_tomador, txt_municipio, txt_sigla_secretaria,
        dsc_situacao_objeto_mcid, dte_inicio_obra_mcid, dte_fim_obra_mcid, vlr_investimento, vlr_repasse, vlr_contrapartida,
        vlr_empenhado, vlr_desembolsado, vlr_pago, prc_execucao_fisica, prc_execucao_financ
        FROM se_saci.view_mat_carteira_investimento WHERE cod_tci = ANY(:t)"""), c, params={'t': tcis}).drop_duplicates('cod_tci').set_index('cod_tci')
    des = pd.read_sql(text("""SELECT cod_tci, CASE WHEN extract(year FROM dte_ordem_bancaria) <= 2022 THEN 'até 2022'
        ELSE extract(year FROM dte_ordem_bancaria)::int::text END AS ano, sum(vlr_ordem_bancaria) AS v
        FROM se_saci.view_mat_desembolsos WHERE cod_tci = ANY(:t) GROUP BY 1, 2"""), c, params={'t': tcis})
dpiv = des.pivot_table(index='cod_tci', columns='ano', values='v', aggfunc='sum').fillna(0)
ANOS_D = ['até 2022', '2023', '2024', '2025', '2026']
for y in ANOS_D:
    if y not in dpiv.columns:
        dpiv[y] = 0.0

sx = s.set_index('cod_tci')
ini_tci = set(g734.cod_tci)


def linha_ct(t, grupo, lic=None):
    d = db.loc[t] if t in db.index else pd.Series(dtype=object); sr = sx.loc[t]
    g = lambda k: d.get(k) if k in d.index else None
    rep_, des_ = g('vlr_repasse'), g('vlr_desembolsado')
    row = {'Grupo': grupo, 'TCI': t, 'Nº TC / convênio': sr.num_convenio, 'Nº proposta (SACI)': g('num_proposta'),
           'Proposta de seleção PAC': sr.num_proposta_selecao_pac, 'Operação Caixa': g('cod_operacao'), 'ID Governa': sr.get('id_governa'),
           'Origem': 'migrada' if sr.kind == 0 else 'nova seleção', 'Fase PAC': sr.dsc_fase_pac, 'Fonte': sr.txt_fonte,
           'Secretaria': g('txt_sigla_secretaria'), 'UF': sr.txt_uf, 'Município': g('txt_municipio') or sr.txt_municipio,
           'Tomador': g('txt_tomador'), 'Modalidade': sr['mod'], 'Ano de seleção': sr.ano or None,
           'Situação do contrato': sr.dsc_situacao_contrato_mcid, 'Situação da obra (SACI)': g('dsc_situacao_objeto_mcid'),
           'Assinatura': dt(sr.dte_assinatura_contrato), 'Início da obra': dt(g('dte_inicio_obra_mcid')),
           'Fim previsto da obra': dt(g('dte_fim_obra_mcid')), '% físico': g('prc_execucao_fisica'), '% financeiro': g('prc_execucao_financ'),
           'Investimento (R$)': g('vlr_investimento'), 'Repasse (R$)': rep_, 'Contrapartida (R$)': g('vlr_contrapartida'),
           'Empenhado (R$)': g('vlr_empenhado'), 'Desembolsado (R$)': des_, 'Pago (R$)': g('vlr_pago'),
           'Saldo a desembolsar (R$)': (rep_ or 0) - (des_ or 0) if rep_ is not None else None,
           'Valor GM migrada pós-2023 (R$)': sr.vlr_gm if sr.kind == 0 else None,
           'Também em obras iniciadas': 'sim' if (lic is not None and t in ini_tci) else None}
    for y in ANOS_D:
        row[f'Desembolso {y} (R$)'] = float(dpiv.loc[t, y]) if t in dpiv.index else 0.0
    if lic is not None:
        row.update({'PC 72 (simulação)': 'sim' if lic.pc72 else 'não', '1ª LAE': dt(lic.dte_primeira_data_lae_tdb),
                    'Publicação do edital': dt(lic.dte_publicacao_licitacao_tgov), 'Prazo legal': lic.legal, 'Prazo efetivo': lic.efetivo,
                    'Situação do prazo': 'vencido' if lic.efetivo < REF else ('até 30/10/2026' if lic.efetivo <= LIM else 'depois de 30/10/2026')})
    return row


L155 = [linha_ct(x.cod_tci_saci, 'Sem edital publicado', x) for _, x in g155.iterrows()]
L257 = [linha_ct(x.cod_tci_saci, 'Edital publicado, sem homologação', x) for _, x in g257.iterrows()]
SN = {1: 'em execução', 2: 'paralisada', 3: 'concluída', 4: 'inacabada'}
L734 = []
for _, x in g734.iterrows():
    row = linha_ct(x.cod_tci, 'Obra iniciada'); row['Situação da obra (SACI + balanço)'] = SN[x.sit_eff]; L734.append(row)
bal = dados.abrir(ext['arquivos']['balanco_governa']).drop_duplicates('id_governa').set_index('id_governa')
for _, x in bxo.iterrows():
    L734.append({'Grupo': 'Obra iniciada', 'TCI': None, 'ID Governa': int(x.id_governa), 'Origem': 'item do balanço sem contrato no SACI',
                 'Fonte': x.fonte, 'Secretaria': x.secretaria, 'UF': x.uf or None, 'Modalidade': x['mod'],
                 'Situação da obra (SACI + balanço)': x.fase, '% físico': x['exec'], 'Investimento (R$)': x.vlr,
                 'Sistema de referência (balanço)': x.sistema_de_referencia,
                 'Código no sistema (balanço)': bal.codigo_do_sistema.get(x.id_governa) if 'codigo_do_sistema' in bal.columns else None})

# Base única: um registro por contrato — quem está em obras iniciadas fica só lá (não soma valor duas vezes)
BASE = L734 + [rw for rw in L155 + L257 if rw['TCI'] not in ini_tci]

COLS = ['Grupo', 'TCI', 'Nº TC / convênio', 'Nº proposta (SACI)', 'Proposta de seleção PAC', 'Operação Caixa', 'ID Governa', 'Origem',
        'Fase PAC', 'Fonte', 'Secretaria', 'UF', 'Município', 'Tomador', 'Modalidade', 'Ano de seleção', 'Situação do contrato',
        'Situação da obra (SACI)', 'Situação da obra (SACI + balanço)', 'Assinatura', 'Início da obra', 'Fim previsto da obra',
        '% físico', '% financeiro', 'Investimento (R$)', 'Repasse (R$)', 'Contrapartida (R$)', 'Empenhado (R$)', 'Desembolsado (R$)',
        'Pago (R$)', 'Saldo a desembolsar (R$)', 'Valor GM migrada pós-2023 (R$)'] + [f'Desembolso {y} (R$)' for y in ANOS_D] + \
       ['PC 72 (simulação)', '1ª LAE', 'Publicação do edital', 'Prazo legal', 'Prazo efetivo', 'Situação do prazo',
        'Também em obras iniciadas', 'Sistema de referência (balanço)', 'Código no sistema (balanço)', 'Link SACI']
wb = Workbook(); fino = Side(style='thin', color='D0D5DD')


def aba(nome, linhas, fora=()):
    cols = [c for c in COLS if c not in fora]
    ws = wb.create_sheet(nome); ws.append(cols)
    for i, h in enumerate(cols, 1):
        cel = ws.cell(row=1, column=i); cel.font = Font(bold=True, color='FFFFFF'); cel.fill = PatternFill('solid', fgColor='071D41')
        cel.alignment = Alignment(wrap_text=True, vertical='center')
        ws.column_dimensions[get_column_letter(i)].width = 15 if '(R$)' in h else 22 if h in (
            'Modalidade', 'Tomador', 'Situação do contrato', 'Grupo', 'Origem') else 13
    ws.row_dimensions[1].height = 44; ws.freeze_panes = 'C2'
    for row in linhas:
        vals = []
        for c in cols:
            v = None if c == 'Link SACI' else row.get(c)
            if v is not None and not isinstance(v, str) and pd.isna(v):
                v = None
            if isinstance(v, pd.Timestamp):
                v = v.date()
            vals.append(v)
        ws.append(vals)
        if row.get('TCI') and 'Link SACI' in cols:
            cel = ws.cell(row=ws.max_row, column=cols.index('Link SACI') + 1); cel.value = 'abrir'
            cel.hyperlink = f"https://saci.cidades.gov.br/contratos/{row['TCI']}"; cel.font = Font(color='1351B4', underline='single')
    for rr in ws.iter_rows(min_row=2, max_col=len(cols)):
        for cel, h in zip(rr, cols):
            cel.border = Border(top=fino, bottom=fino, left=fino, right=fino)
            if '(R$)' in h:
                cel.number_format = '#,##0.00'
            elif h.startswith('%'):
                cel.number_format = '0.0'
            elif hasattr(cel.value, 'year'):
                cel.number_format = 'DD/MM/YYYY'
    ws.auto_filter.ref = ws.dimensions


LIC = ['PC 72 (simulação)', '1ª LAE', 'Publicação do edital', 'Prazo legal', 'Prazo efetivo', 'Situação do prazo', 'Também em obras iniciadas']
BAL = ['Sistema de referência (balanço)', 'Código no sistema (balanço)', 'Situação da obra (SACI + balanço)']
aba(f'Sem edital ({len(L155)})', L155, BAL)
aba(f'Edital sem homologação ({len(L257)})', L257, BAL)
aba(f'Obras iniciadas ({len(L734)})', L734, LIC)
aba(f'Base única ({len(BASE)})', BASE)

rs = wb['Sheet']; rs.title = 'Resumo'
tot = lambda L, k: float(sum((rw.get(k) or 0) for rw in L))
sob155 = sum(1 for rw in L155 if rw['TCI'] in ini_tci); sob257 = sum(1 for rw in L257 if rw['TCI'] in ini_tci)
lin = [['Contratos OGU para a previsão orçamentária — apresentação Novo PAC (MCid)'], [],
       ['Extração', ext['extraido_em'] + ' (SACI ' + ext['atualizacao']['saci'] + ' · painelpac · balanço Governa 30/06/2026)'], [],
       ['Grupo', 'Contratos', 'Repasse (R$)', 'Desembolsado (R$)', 'Saldo a desembolsar (R$)']]
for nome, L in [('Sem edital publicado (slide 3.2 — publicação do edital)', L155),
                ('Edital publicado, sem homologação (slide 3.2 — fim da licitação)', L257),
                ('Obras iniciadas OGU (slide 5.2 — fonte OGU, migradas + novas, SACI + balanço)', L734),
                ('Base única (sem repetir contrato)', BASE)]:
    lin.append([nome, len(L), tot(L, 'Repasse (R$)'), tot(L, 'Desembolsado (R$)'), tot(L, 'Saldo a desembolsar (R$)')])
lin += [[], ['Critérios'],
        ['Sem edital / sem homologação', 'OGU do Novo PAC Seleção em "Contratado - Normal" (painelpac), sem bln_pac = NÃO e sem os sem modalidade; '
         'inclui os assinados antes de 21/10/2025 (PC 72) em simulação, pela mesma regra (LAE + 60 dias / publicação + 120 dias).'],
        ['Prazo efetivo', 'o maior entre o prazo legal e 30/10/2026; situação do prazo na data da extração.'],
        ['Obras iniciadas', 'OGU assinados (sem em contratação e cancelados), Novo PAC Seleção + Migrado, com obra em execução, paralisada, concluída '
         'ou inacabada pelo SACI ou pelo balanço jun/26; + ' + str(len(bxo)) + ' itens do balanço sem contrato no SACI (sem TCI: ID Governa e código).'],
        ['Sobreposição', f'{sob155} sem edital e {sob257} sem homologação já têm obra iniciada (coluna "Também em obras iniciadas"); '
         'na Base única ficam só em obras iniciadas, para não somar o valor duas vezes.'],
        ['Saldo a desembolsar', 'repasse − desembolsado do SACI (as migradas trazem também o valor GM pós-2023).'],
        ['Desembolso por ano', 'soma das ordens bancárias (se_saci.view_mat_desembolsos) por TCI e ano.']]
for row in lin:
    rs.append(row)
rs['A1'].font = Font(bold=True, size=14, color='071D41'); rs.column_dimensions['A'].width = 72
for col in 'BCDE':
    rs.column_dimensions[col].width = 22
for rr in rs.iter_rows(min_row=6, max_row=9):
    for cel in rr[2:5]:
        cel.number_format = '#,##0.00'
for cel in rs[5]:
    cel.font = Font(bold=True)
os.makedirs(os.path.dirname(SAIDA), exist_ok=True)
wb.save(SAIDA)
print('escrito:', SAIDA)
for nome, L in [('sem edital', L155), ('sem homologação', L257), ('obras iniciadas', L734), ('base única', BASE)]:
    ts = [rw['TCI'] for rw in L if rw['TCI']]
    print(f'  {nome}: {len(L)} | sem TCI {sum(1 for rw in L if not rw["TCI"])} | TCI repetido {len(ts) - len(set(ts))} | '
          f'repasse R$ {tot(L, "Repasse (R$)") / 1e9:.2f} bi | saldo R$ {tot(L, "Saldo a desembolsar (R$)") / 1e9:.2f} bi')
